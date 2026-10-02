"""
Serial COM3 Reader → Web Dashboard con Excel Export
====================================================
Servidor FastAPI que lee datos CSV del puerto COM3,
los envía al frontend vía WebSocket en tiempo real,
y permite descargar todos los datos como archivo Excel.
"""

import asyncio
import json
import os
import time
from datetime import datetime
from io import BytesIO
from typing import List, Optional

import serial
import serial.tools.list_ports
from fastapi import FastAPI, WebSocket, WebSocketDisconnect
from fastapi.responses import FileResponse, HTMLResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ─── Configuración ──────────────────────────────────────────────
SERIAL_PORT = "COM3"
BAUD_RATE = 115200
# Nombres de las columnas de datos del monocóptero
COLUMN_NAMES: List[str] = ["Tiempo Relativo (s)", "PWM", "Altura (cm)"]

# Modo demo: simula datos si el puerto COM3 no está disponible
DEMO_MODE = False  # Se activa automáticamente si COM3 no se encuentra

# --- Almacenamiento de datos ----
data_store: List[dict] = []
send_log: List[dict] = []  # Historial de comandos enviados
num_columns: Optional[int] = None
column_names: List[str] = []

# Referencia global a la conexion serial (para enviar datos)
serial_connection: Optional[serial.Serial] = None

# ─── FastAPI App ────────────────────────────────────────────────
app = FastAPI(title="Serial COM3 Reader")

# Servir archivos estáticos (HTML/CSS/JS)
static_dir = os.path.join(os.path.dirname(__file__), "static")
os.makedirs(static_dir, exist_ok=True)
app.mount("/static", StaticFiles(directory=static_dir), name="static")

# Clientes WebSocket conectados
connected_clients: List[WebSocket] = []


def detect_serial_port() -> Optional[str]:
    """Detecta si COM3 está disponible."""
    ports = serial.tools.list_ports.comports()
    for port in ports:
        if port.device == SERIAL_PORT:
            return port.device
    return None


def parse_csv_line(line: str) -> Optional[List[float]]:
    """Parsea una línea CSV y retorna lista de floats."""
    global num_columns, column_names
    line = line.strip()
    if not line:
        return None

    parts = line.split(",")
    values = []
    for part in parts:
        part = part.strip()
        try:
            values.append(float(part))
        except ValueError:
            return None  # Línea no numérica, ignorar

    if not values:
        return None

    # Auto-detectar número de columnas en la primera lectura válida
    if num_columns is None:
        num_columns = len(values)
        if COLUMN_NAMES and len(COLUMN_NAMES) == num_columns:
            column_names = COLUMN_NAMES
        else:
            column_names = [f"Sensor {i+1}" for i in range(num_columns)]
        print(f"[OK] Detectadas {num_columns} columnas: {column_names}")

    # Validar que la línea tenga el número esperado de columnas
    if len(values) != num_columns:
        return None

    return values


async def broadcast(message: dict):
    """Envía un mensaje a todos los clientes WebSocket conectados."""
    disconnected = []
    msg_text = json.dumps(message)
    for client in connected_clients:
        try:
            await client.send_text(msg_text)
        except Exception:
            disconnected.append(client)
    for client in disconnected:
        connected_clients.remove(client)


async def read_serial():
    """Lee datos del puerto serial en un loop asincrono."""
    global DEMO_MODE, serial_connection

    # Intentar abrir puerto serial directamente
    ser = None

    try:
        ser = serial.Serial(SERIAL_PORT, BAUD_RATE, timeout=1)
        serial_connection = ser  # Guardar referencia global
        print(f"[OK] Conectado a {SERIAL_PORT} @ {BAUD_RATE} baud", flush=True)

        # Esperar a que Arduino termine de reiniciarse tras la conexion serial
        print("[...] Esperando 2s para que Arduino se reinicie...", flush=True)
        await asyncio.sleep(2)
        # Limpiar buffer de datos basura del reset
        ser.reset_input_buffer()
        print("[OK] Arduino listo, leyendo datos...", flush=True)

        await broadcast({
            "type": "status",
            "connected": True,
            "port": SERIAL_PORT,
            "baudrate": BAUD_RATE
        })
    except serial.SerialException as e:
        print(f"[ERROR] No se pudo abrir {SERIAL_PORT}: {e}", flush=True)
        print("[!] Activando modo DEMO.", flush=True)
        DEMO_MODE = True
        ser = None

    if DEMO_MODE:
        print("[DEMO] Modo DEMO activo - generando datos simulados", flush=True)
        await broadcast({
            "type": "status",
            "connected": True,
            "port": "DEMO",
            "baudrate": BAUD_RATE,
            "demo": True
        })
        await demo_data_generator()
        return

    # Loop de lectura serial real usando run_in_executor para no bloquear
    loop = asyncio.get_event_loop()

    def _blocking_readline():
        """Lee una linea del serial (bloqueante, corre en thread)."""
        try:
            if ser and ser.is_open and ser.in_waiting > 0:
                return ser.readline().decode("utf-8", errors="ignore")
        except Exception:
            pass
        return None

    try:
        while True:
            # Ejecutar lectura bloqueante en un thread separado
            raw_line = await loop.run_in_executor(None, _blocking_readline)

            if raw_line:
                values = parse_csv_line(raw_line)
                if values is not None:
                    now = datetime.now()
                    entry = {
                        "timestamp": now.strftime("%H:%M:%S.%f")[:-3],
                        "datetime": now.isoformat(),
                        "values": values,
                        "raw": raw_line.strip()
                    }
                    data_store.append(entry)

                    await broadcast({
                        "type": "data",
                        "timestamp": entry["timestamp"],
                        "values": values,
                        "index": len(data_store) - 1,
                        "columns": column_names
                    })

            await asyncio.sleep(0.01)  # 10ms entre polls

    except asyncio.CancelledError:
        pass
    finally:
        if ser:
            ser.close()
            print("[OK] Puerto serial cerrado", flush=True)


async def demo_data_generator():
    """Genera datos simulados de vuelo de monocoptero."""
    import math
    import random

    global num_columns, column_names

    num_columns = 3
    column_names = COLUMN_NAMES

    # Enviar nombres de columnas
    await broadcast({
        "type": "config",
        "columns": column_names,
        "num_columns": num_columns
    })

    t = 0
    tiempo_rel = 0.0
    altura = 0.0
    velocidad = 0.0

    try:
        while True:
            # Simular perfil de vuelo de monocoptero
            # Fase 1: Reposo (0-3s)
            # Fase 2: Rampa de subida PWM (3-8s)
            # Fase 3: Hover estable (8-18s)
            # Fase 4: Descenso (18-23s)
            # Fase 5: Reposo (23-26s), luego repite
            cycle_time = tiempo_rel % 26.0

            if cycle_time < 3.0:
                # Reposo
                pwm = 1000 + random.gauss(0, 2)
                target_alt = 0.0
            elif cycle_time < 8.0:
                # Rampa de subida
                progress = (cycle_time - 3.0) / 5.0
                pwm = 1000 + 600 * progress + random.gauss(0, 5)
                target_alt = 80.0 * progress ** 1.5
            elif cycle_time < 18.0:
                # Hover estable con oscilaciones
                pwm = 1550 + 50 * math.sin(cycle_time * 0.8) + random.gauss(0, 8)
                target_alt = 75.0 + 8.0 * math.sin(cycle_time * 0.5) + random.gauss(0, 1.5)
            elif cycle_time < 23.0:
                # Descenso
                progress = (cycle_time - 18.0) / 5.0
                pwm = 1550 - 500 * progress + random.gauss(0, 5)
                target_alt = 75.0 * (1.0 - progress ** 1.3)
            else:
                # Reposo final
                pwm = 1000 + random.gauss(0, 2)
                target_alt = 0.0

            # Simular inercia en la altura
            altura += (target_alt - altura) * 0.15 + random.gauss(0, 0.3)
            altura = max(0.0, altura)

            values = [
                round(tiempo_rel, 2),
                round(max(1000, min(2000, pwm)), 0),
                round(altura, 1)
            ]
            now = datetime.now()

            entry = {
                "timestamp": now.strftime("%H:%M:%S.%f")[:-3],
                "datetime": now.isoformat(),
                "values": values,
                "raw": ",".join(str(v) for v in values)
            }
            data_store.append(entry)

            await broadcast({
                "type": "data",
                "timestamp": entry["timestamp"],
                "values": values,
                "index": len(data_store) - 1,
                "columns": column_names
            })

            t += 1
            tiempo_rel += 0.2
            await asyncio.sleep(0.2)  # 5 lecturas/segundo en demo

    except asyncio.CancelledError:
        pass


# ─── Eventos de inicio ─────────────────────────────────────────
serial_task = None


@app.on_event("startup")
async def startup():
    global serial_task
    serial_task = asyncio.create_task(read_serial())


@app.on_event("shutdown")
async def shutdown():
    if serial_task:
        serial_task.cancel()
        try:
            await serial_task
        except asyncio.CancelledError:
            pass


# ─── Rutas ──────────────────────────────────────────────────────
@app.get("/")
async def root():
    """Sirve la página principal."""
    index_path = os.path.join(static_dir, "index.html")
    return FileResponse(index_path)


@app.websocket("/ws")
async def websocket_endpoint(websocket: WebSocket):
    """WebSocket para datos en tiempo real."""
    await websocket.accept()
    connected_clients.append(websocket)
    print(f"[WS] Cliente WebSocket conectado ({len(connected_clients)} total)")

    # Enviar configuración actual si ya se detectaron columnas
    if column_names:
        await websocket.send_text(json.dumps({
            "type": "config",
            "columns": column_names,
            "num_columns": num_columns
        }))

    # Enviar historial existente
    if data_store:
        history = []
        for entry in data_store[-200:]:  # Últimos 200 puntos
            history.append({
                "timestamp": entry["timestamp"],
                "values": entry["values"]
            })
        await websocket.send_text(json.dumps({
            "type": "history",
            "data": history,
            "columns": column_names
        }))

    try:
        while True:
            # Mantener conexion abierta, escuchar comandos
            msg = await websocket.receive_text()
            cmd = json.loads(msg)
            if cmd.get("action") == "clear":
                data_store.clear()
                send_log.clear()
                await broadcast({"type": "cleared"})
                print("[OK] Datos limpiados")
            elif cmd.get("action") == "send":
                value = cmd.get("value", "").strip()
                if value:
                    now = datetime.now()
                    sent_entry = {
                        "timestamp": now.strftime("%H:%M:%S.%f")[:-3],
                        "value": value
                    }
                    send_log.append(sent_entry)

                    if serial_connection and serial_connection.is_open:
                        try:
                            serial_connection.write((value + "\n").encode("utf-8"))
                            print(f"[TX] Enviado a {SERIAL_PORT}: {value}")
                            await broadcast({
                                "type": "sent",
                                "timestamp": sent_entry["timestamp"],
                                "value": value,
                                "status": "ok"
                            })
                        except Exception as e:
                            print(f"[ERROR] Error enviando: {e}")
                            await broadcast({
                                "type": "sent",
                                "timestamp": sent_entry["timestamp"],
                                "value": value,
                                "status": "error",
                                "error": str(e)
                            })
                    else:
                        # Modo demo - simular envio
                        print(f"[TX-DEMO] Enviado (simulado): {value}")
                        await broadcast({
                            "type": "sent",
                            "timestamp": sent_entry["timestamp"],
                            "value": value,
                            "status": "demo"
                        })
    except WebSocketDisconnect:
        connected_clients.remove(websocket)
        print(f"[WS] Cliente desconectado ({len(connected_clients)} restantes)")
    except Exception:
        if websocket in connected_clients:
            connected_clients.remove(websocket)


@app.get("/download-excel")
async def download_excel():
    """Genera y descarga un archivo Excel con todos los datos."""
    if not data_store:
        return HTMLResponse("<h3>No hay datos para descargar</h3>", status_code=400)

    wb = Workbook()
    ws = wb.active
    ws.title = "Datos Serial COM3"

    # ─── Estilos ────────────────────────────────────────────────
    header_font = Font(name="Calibri", bold=True, color="FFFFFF", size=12)
    header_fill = PatternFill(start_color="1a1a2e", end_color="1a1a2e", fill_type="solid")
    header_border = Border(
        bottom=Side(style="medium", color="6c63ff")
    )
    header_align = Alignment(horizontal="center", vertical="center")

    data_font = Font(name="Calibri", size=11)
    data_align = Alignment(horizontal="center", vertical="center")
    alt_fill = PatternFill(start_color="F5F5FF", end_color="F5F5FF", fill_type="solid")

    # ─── Encabezados ────────────────────────────────────────────
    cols = column_names if column_names else [f"Sensor {i+1}" for i in range(num_columns or 1)]
    headers = ["#", "Timestamp"] + cols

    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = header_border
        cell.alignment = header_align

    # ─── Datos ──────────────────────────────────────────────────
    for row_idx, entry in enumerate(data_store, 2):
        ws.cell(row=row_idx, column=1, value=row_idx - 1).font = data_font
        ws.cell(row=row_idx, column=1, value=row_idx - 1).alignment = data_align

        ts_cell = ws.cell(row=row_idx, column=2, value=entry["timestamp"])
        ts_cell.font = data_font
        ts_cell.alignment = data_align

        for val_idx, val in enumerate(entry["values"]):
            cell = ws.cell(row=row_idx, column=3 + val_idx, value=val)
            cell.font = data_font
            cell.alignment = data_align
            cell.number_format = "0.00"

        # Filas alternadas
        if row_idx % 2 == 0:
            for col_idx in range(1, len(headers) + 1):
                ws.cell(row=row_idx, column=col_idx).fill = alt_fill

    # ─── Ancho de columnas ──────────────────────────────────────
    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 16
    for i in range(len(cols)):
        col_letter = get_column_letter(3 + i)
        ws.column_dimensions[col_letter].width = 18

    # ─── Guardar en buffer ──────────────────────────────────────
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    filename = f"datos_serial_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )


@app.get("/api/status")
async def api_status():
    """Retorna el estado actual del sistema."""
    return {
        "serial_port": SERIAL_PORT,
        "baudrate": BAUD_RATE,
        "demo_mode": DEMO_MODE,
        "data_points": len(data_store),
        "columns": column_names,
        "connected_clients": len(connected_clients)
    }


@app.get("/api/ports")
async def list_ports():
    """Lista los puertos seriales disponibles."""
    ports = serial.tools.list_ports.comports()
    return [
        {"device": p.device, "description": p.description}
        for p in ports
    ]


# ─── Punto de entrada ──────────────────────────────────────────
if __name__ == "__main__":
    import uvicorn
    print("=" * 55)
    print("  Serial COM3 Reader - Web Dashboard")
    print("  Abre http://localhost:8000 en tu navegador")
    print("=" * 55)
    uvicorn.run(app, host="0.0.0.0", port=8000, log_level="info")
